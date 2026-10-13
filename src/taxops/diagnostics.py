"""Isolated, synthetic checks executed by the packaged EXE as well as source.

This is a service/native-library smoke, not manual GUI or clean-OS acceptance.
No caller-provided data root is accepted. All business data is temporary.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import platform
import sqlite3
import sys
import tempfile
from contextlib import contextmanager
from pathlib import Path

from . import __version__


class SelfTestFailure(RuntimeError):
    pass


def _require(condition: bool, message: str) -> None:
    if not condition:
        raise SelfTestFailure(message)


@contextmanager
def _temporary_profile_paths(root: Path):
    """Keep environment-based runtime caches inside the synthetic data root.

    This is a scoped path redirect, not an OS sandbox. Native registry settings
    or APIs that ignore these environment variables are outside this boundary.
    """
    paths = {
        "LOCALAPPDATA": root / "profile" / "local",
        "APPDATA": root / "profile" / "roaming",
        "USERPROFILE": root / "profile" / "home",
        "HOME": root / "profile" / "home",
        "XDG_CACHE_HOME": root / "profile" / "cache",
        "XDG_CONFIG_HOME": root / "profile" / "config",
        "XDG_DATA_HOME": root / "profile" / "data",
    }
    previous = {key: os.environ.get(key) for key in paths}
    try:
        for key, path in paths.items():
            path.mkdir(parents=True, exist_ok=True)
            os.environ[key] = str(path)
        yield
    finally:
        for key, value in previous.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value


@contextmanager
def _deny_python_network(attempts: list[str]):
    """Block CPython socket audit events; this is not an OS/network firewall."""
    active = True
    denied = {"socket.connect", "socket.getaddrinfo", "socket.gethostbyname",
              "socket.gethostbyaddr", "socket.sendto", "socket.sendmsg", "socket.bind"}

    def audit(event, _args):
        if active and event in denied:
            attempts.append(event)
            raise SelfTestFailure(f"Network operation forbidden during self-test: {event}")

    sys.addaudithook(audit)
    try:
        yield
    finally:
        active = False


def _run_checks(root: Path, report: dict) -> None:
    from PySide6 import __version__ as qt_binding_version
    from PySide6.QtCore import QCoreApplication, QEvent, QSize
    from PySide6.QtGui import QIcon, QImage, QImageReader, QPainter, QPdfWriter
    from PySide6.QtPdf import QPdfDocument
    from PySide6.QtPdfWidgets import QPdfView
    from PySide6.QtWidgets import QApplication
    import openpyxl
    import ssl

    from .core.paths import resolve_paths
    from .core.resources import asset_path
    from .db.migrations import MIGRATIONS
    from .services.attachments import UploadAttachmentInput
    from .services.canvas_notes import CreateCanvasNoteInput
    from .services.clients import CreateClientInput
    from .services.client_leases import LeaseInput
    from .services.compliance_profiles import ComplianceProfileItemInput
    from .services.document_requests import CreateDocumentRequestInput
    from .services.engagements import CreateEngagementInput
    from .services.tasks import CreateTaskInput
    from .services.work_records import (
        CreateWorkflowTemplateInput, WorkflowStageInput, WorkflowStepInput,
    )
    from .ui.app import bootstrap
    from .ui.main_window import MainWindow
    from .ui.style import apply
    from .diagnostics_qt_language import verify_chinese_widgets
    from .diagnostics_notices import verify_notice_widgets

    def passed(name, **detail):
        report["checks"].append({"name": name, "status": "passed", **detail})

    app = QApplication.instance() or QApplication([])
    report.update(qt_binding=qt_binding_version, qt_platform=app.platformName())
    apply(app)
    passed("qt_traditional_chinese_widgets_and_fallback", **verify_chinese_widgets(app, root))
    icon = asset_path("app_icon.ico")
    _require(icon.is_file() and not QIcon(str(icon)).isNull(), "Packaged icon unavailable")
    passed("qt_application_and_icon")
    formats = {bytes(value).decode("ascii").lower() for value in QImageReader.supportedImageFormats()}
    _require({"png", "jpeg", "bmp", "ico"}.issubset(formats), "Required image plugin missing")
    png = root / "合成圖片.png"
    image = QImage(16, 16, QImage.Format.Format_RGB32)
    image.fill(0x336699)
    _require(image.save(str(png)), "PNG writer failed")
    _require(not QImage(str(png)).isNull(), "PNG reader failed")
    image_bytes = png.read_bytes()
    passed("image_codecs", formats=sorted(formats))

    pdf = root / "合成文件.pdf"
    writer = QPdfWriter(str(pdf))
    painter = QPainter(writer)
    _require(painter.isActive(), "PDF writer unavailable")
    painter.drawText(100, 200, "TaxOps synthetic offline PDF")
    painter.end()
    del painter, writer
    document = QPdfDocument()
    try:
        _require(document.load(str(pdf)) == QPdfDocument.Error.None_, "QtPdf cannot load generated PDF")
        _require(document.pageCount() == 1, "Unexpected generated PDF page count")
        _require(not document.render(0, QSize(100, 100)).isNull(), "QtPdf cannot render page")
        view = QPdfView()
        view.setDocument(document)
        view.close()
        view.deleteLater()
        QCoreApplication.sendPostedEvents(view, QEvent.Type.DeferredDelete)
    finally:
        document.close()
        document.deleteLater()
        QCoreApplication.sendPostedEvents(document, QEvent.Type.DeferredDelete)
    passed("qt_pdf_and_pdf_widgets")
    context = ssl.create_default_context()
    _require(context.verify_mode == ssl.CERT_REQUIRED and context.check_hostname,
             "TLS verification defaults are disabled")
    passed("ssl_runtime", openssl=ssl.OPENSSL_VERSION,
           boundary="Import/context only; no live HTTPS or certificate-chain test")

    source = bootstrap(resolve_paths(override_root=root / "來源 資料"))
    try:
        versions = {row["version"] for row in source.conn.execute("SELECT version FROM schema_migrations")}
        _require(versions == {v for v, _ in MIGRATIONS}, "Migration set incomplete")
        passed("all_migrations", count=len(versions), sqlite=sqlite3.sqlite_version)
        window = MainWindow(source)
        try:
            window.show()
            app.processEvents()
            _require(window.isVisible(), "MainWindow failed to become visible")
            passed("offline_third_party_notices_view_and_export",
                   **verify_notice_widgets(window, app, root))
        finally:
            window.close()
            window.deleteLater()
            QCoreApplication.sendPostedEvents(window, QEvent.Type.DeferredDelete)
        passed("main_window_construct_show_close",
               boundary="No visual, interaction, scaling, or long-running acceptance")
        client = source.clients.create_client(CreateClientInput(
            "SMOKE001", "離線驗證客戶", registered_address="合成登記地址",
            contact_address="合成聯絡地址", contact_address_same=False,
        ))
        _require(client.registered_address == "合成登記地址"
                 and client.contact_address == "合成聯絡地址"
                 and client.contact_address_same is False, "Client address model failed")
        lease = source.client_leases.create_lease(client.id, LeaseInput(
            lease_name="合成租約", start_date="2026-01-01", end_date="2026-12-31",
            monthly_rent=10_000,
        ))
        lease_attachment = source.attachments.upload_lease_attachment(client.id, lease.id, pdf)
        source.attachments.accept_attachment(lease_attachment.id)
        _require(source.attachments.resolve_file_path(lease_attachment.id).read_bytes() == pdf.read_bytes(),
                 "Client lease attachment copy differs")
        passed("client_addresses_lease_and_lease_attachment")
        case = source.engagements.create_engagement(
            CreateEngagementInput(client.id, "合成營業稅案件", "vat", "2026-10")
        )
        request, items = source.doc_requests.create_request(CreateDocumentRequestInput(
            case.id, "vat", "2026-10", item_names=("合成進項憑證", "合成銀行明細"),
        ))
        _require(len(items) == 2, "Document items not persisted")
        task = source.tasks.create_task(CreateTaskInput(case.id, "合成核對工作", client_id=client.id))
        _require(task.id > 0, "Task not persisted")
        _require(bool(source.search.search_clients("離線驗證")), "SQLite FTS client query failed")
        passed("client_case_document_task_and_fts")
        source.compliance_profiles.upsert_profile(
            client.id, fiscal_year_start_month=1,
            items=(ComplianceProfileItemInput("vat", "bimonthly"),),
        )
        drafts = source.annual_work.preview(client.id, 2026)
        generated = source.annual_work.confirm_preview(client.id, 2026, drafts)
        repeated = source.annual_work.confirm_preview(client.id, 2026, drafts)
        _require(len(generated.items) == 6 and repeated.inserted_item_count == 0,
                 "Annual generation or repeat confirmation failed")
        annual_item = generated.items[0]
        source.annual_work.link_existing_engagement(annual_item.id, case.id)
        annual_task = source.annual_work.create_linked_task(annual_item.id, title="合成年度核對")
        for category, amount in (("tax_liability", 1_000), ("tax_payment", 300),
                                 ("fee_receivable", 200), ("fee_receipt", 50)):
            source.annual_transactions.add(annual_item.id, category, amount, "2026-03-01")
        balance = source.annual_transactions.balance(annual_item.id)
        _require(balance.unpaid_tax == 700 and balance.outstanding_fee == 150,
                 "Annual tax and fee balances differ")
        passed("annual_profile_generation_linked_task_and_transactions")
        attachment = source.attachments.upload_attachment(UploadAttachmentInput(case.id, request.id, pdf))
        source.attachments.accept_attachment(attachment.id)
        pdf_bytes = pdf.read_bytes()
        _require(source.attachments.resolve_file_path(attachment.id).read_bytes() == pdf_bytes,
                 "Attachment copy differs")
        passed("attachment_upload_accept_readback", sha256=hashlib.sha256(pdf_bytes).hexdigest())
        template = source.work_records.create_template(CreateWorkflowTemplateInput(
            "合成工作流程", (WorkflowStageInput("核對", (WorkflowStepInput("完成核對"),)),),
        ))
        source.work_records.set_template_image_path(template.id, str(png))
        source.work_records.instantiate_run(template.id)
        note = source.canvas_notes.create_note(CreateCanvasNoteInput("合成圖像筆記"))
        asset = source.canvas_notes.import_image_asset(png)
        scene = json.loads(note.scene_json)
        scene["objects"].append({"type": "image", "asset_path": asset})
        source.canvas_notes.update_note(note.id, title=note.title, scene_json=json.dumps(scene))
        passed("workflow_and_canvas_assets")
        xlsx = root / "合成匯出.xlsx"
        count = source.export.export_missing_items_xlsx(xlsx)
        book = openpyxl.load_workbook(xlsx, read_only=True)
        try:
            _require(count == 2 and book.active.max_row == 3, "XLSX export/readback failed")
        finally:
            book.close()
        passed("xlsx_roundtrip", rows=count)
        backup = Path(source.backup.create_backup(source.paths).backup_path)
        _require(backup.name.endswith(".taxops.zip"), "Full portable backup format unavailable")
        passed("full_backup", bytes=backup.stat().st_size)
    finally:
        source.close()

    target_paths = resolve_paths(override_root=root / "還原 資料")
    target = bootstrap(target_paths)
    try:
        target.backup.restore_backup(backup, target_paths)
    finally:
        target.close()
    target = bootstrap(target_paths)
    try:
        restored_client = target.clients.get_client(client.id)
        _require(restored_client is not None and restored_client.client_name == "離線驗證客戶",
                 "Restored client missing")
        _require(restored_client.registered_address == "合成登記地址"
                 and restored_client.contact_address == "合成聯絡地址"
                 and restored_client.contact_address_same is False, "Restored addresses differ")
        restored_lease = target.client_leases.get_lease(lease.id)
        _require(restored_lease is not None and restored_lease.monthly_rent == 10_000
                 and restored_lease.client_id == client.id, "Restored lease differs")
        lease_row = target.attachments.get(lease_attachment.id)
        _require(lease_row is not None and lease_row.client_id == client.id
                 and lease_row.lease_id == lease.id and lease_row.engagement_id is None
                 and lease_row.status == "accepted", "Restored lease attachment ownership differs")
        _require(target.attachments.resolve_file_path(lease_attachment.id).read_bytes() == pdf_bytes,
                 "Restored lease attachment bytes differ")
        _require(target.conn.execute("SELECT COUNT(*) FROM engagements").fetchone()[0] == 1,
                 "Restored case missing")
        _require(target.conn.execute("SELECT COUNT(*) FROM document_request_items").fetchone()[0] == 2,
                 "Restored document items missing")
        _require(target.conn.execute("SELECT COUNT(*) FROM workflow_tasks").fetchone()[0] == 2,
                 "Restored task missing")
        _require(target.attachments.get(attachment.id).status == "accepted", "Attachment status not restored")
        _require(target.attachments.resolve_file_path(attachment.id).read_bytes() == pdf_bytes,
                 "Restored attachment differs")
        for record in (target.work_records.list_templates()[0], target.work_records.list_runs()[0]):
            relative = json.loads(record.context_snapshot)["image_path"]
            _require((target.work_records.workflow_assets_dir / relative).read_bytes() == image_bytes,
                     "Restored workflow image differs")
        restored_note = target.canvas_notes.get_note(note.id)
        relative = json.loads(restored_note.scene_json)["objects"][0]["asset_path"]
        _require((target.canvas_notes.note_assets_dir / relative).read_bytes() == image_bytes,
                 "Restored canvas image differs")
        profile = target.conn.execute(
            "SELECT fiscal_year_start_month FROM compliance_profiles WHERE client_id = ?", (client.id,)
        ).fetchone()
        _require(profile is not None and profile[0] == 1, "Restored annual profile missing")
        restored_items = target.conn.execute(
            "SELECT id FROM annual_work_items WHERE workspace_id = ? ORDER BY id", (generated.workspace.id,)
        ).fetchall()
        _require([row[0] for row in restored_items] == sorted(item.id for item in generated.items),
                 "Restored annual work set differs")
        restored_task = target.tasks.get_task(annual_task.id)
        _require(restored_task is not None and restored_task.annual_work_item_id == annual_item.id,
                 "Restored annual task link missing")
        restored_balance = target.annual_transactions.balance(annual_item.id)
        _require(restored_balance.unpaid_tax == 700 and restored_balance.outstanding_fee == 150,
                 "Restored annual transactions differ")
        _require(len(target.annual_transactions.list_for_work_item(annual_item.id)) == 4,
                 "Restored annual transaction events differ")
        passed("client_lease_and_annual_full_restore")
        _require(bool(target.search.search_clients("離線驗證")), "Restored FTS index failed")
        _require(target.conn.execute("PRAGMA integrity_check").fetchone()[0] == "ok", "Integrity check failed")
        _require(not target.conn.execute("PRAGMA foreign_key_check").fetchall(), "Foreign keys invalid")
        actions = {r[0] for r in target.conn.execute("SELECT action FROM audit_logs")}
        _require({"client.create", "engagement.create", "doc_request.create", "attachment.upload",
                  "attachment.accept", "backup.restore", "client.lease.create",
                  "compliance_profile.update", "annual_workspace.confirm",
                  "annual_transaction.add", "annual_work.task.create"}.issubset(actions), "Expected audit trail missing")
        passed("full_restore_reopen_assets_fts_and_audit")
    finally:
        target.close()


def run_self_test(report_path: Path, *, require_frozen: bool = False) -> int:
    report = {"schema": 1, "status": "failed", "app_version": __version__,
              "frozen": bool(getattr(sys, "frozen", False)), "platform": sys.platform,
              "python": platform.python_version(), "checks": [], "network_attempts": [],
              "network_guard": "CPython socket audit events only; not an OS firewall",
              "profile_isolation": "Temporary environment paths only; not a native registry or OS sandbox",
              "clean_os_verified": False, "manual_gui_verified": False}
    try:
        _require(not require_frozen or report["frozen"], "This check requires the frozen EXE")
        with tempfile.TemporaryDirectory(prefix="taxops_self_test_") as temp:
            with _temporary_profile_paths(Path(temp)):
                with _deny_python_network(report["network_attempts"]):
                    _run_checks(Path(temp), report)
            _require(not report["network_attempts"], "Self-test attempted network access")
        report["checks"].append({"name": "temporary_data_cleanup", "status": "passed"})
        if report["frozen"]:
            from build_tools.windows_loaded_modules import capture_loaded_modules
            report["loaded_modules"] = capture_loaded_modules("after_18_checks")
            _require(report["loaded_modules"]["status"] == "observed",
                     "Frozen diagnostics require actual Windows loaded-module evidence")
        else:
            report["loaded_modules"] = {"status": "unsupported", "reason": "Source execution is not frozen Windows module evidence"}
        report["status"] = "passed"
    except Exception as exc:
        report["error"] = {"type": type(exc).__name__, "message": str(exc)}
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return 0 if report["status"] == "passed" else 1


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument("--require-frozen", action="store_true")
    args = parser.parse_args(argv)
    return run_self_test(args.report, require_frozen=args.require_frozen)


if __name__ == "__main__":
    raise SystemExit(main())
